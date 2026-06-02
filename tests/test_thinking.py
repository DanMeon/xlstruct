"""ExtractionEngine: thinking unification (C1), confidence fail-fast (B2), error scope (C4)."""

import io
from unittest.mock import AsyncMock, MagicMock, patch

import openpyxl
import pytest
from pydantic import BaseModel

from xlstruct.config import ExtractorConfig, thinking_call_kwargs
from xlstruct.exceptions import ErrorCode, ExtractionError
from xlstruct.extraction.engine import ExtractionEngine, _split_confidence
from xlstruct.extractor import Extractor
from xlstruct.schemas.suggest import FieldDef, SuggestedFields

THINKING_BLOCK = {"type": "enabled", "budget_tokens": 10_000}


class Sample(BaseModel):
    name: str
    value: int


def _mock_client(return_value: object) -> MagicMock:
    # ^ usage=None → UsageTracker.record() treats it as zero usage and no-ops
    completion = MagicMock(usage=None)
    client = MagicMock()
    client.create_with_completion = AsyncMock(return_value=(return_value, completion))
    return client


# * thinking_call_kwargs


class TestThinkingCallKwargs:
    def test_empty_when_thinking_off(self):
        cfg = ExtractorConfig(provider="anthropic/claude-sonnet-4-6", thinking=False)
        assert thinking_call_kwargs(cfg) == {}

    def test_empty_for_non_anthropic_even_when_on(self):
        cfg = ExtractorConfig(provider="openai/gpt-4o", thinking=True)
        assert thinking_call_kwargs(cfg) == {}

    def test_block_for_anthropic_thinking(self):
        cfg = ExtractorConfig(provider="anthropic/claude-sonnet-4-6", thinking=True)
        kw = thinking_call_kwargs(cfg)
        assert kw["temperature"] == 1
        assert kw["thinking"] == THINKING_BLOCK
        assert kw["max_tokens"] == 16_000
        assert kw["model"] == "claude-sonnet-4-6"


# * Direct extraction honors thinking (the C1 bug)


class TestDirectExtractionThinking:
    async def test_thinking_kwargs_reach_llm(self):
        cfg = ExtractorConfig(provider="anthropic/claude-sonnet-4-6", thinking=True)
        client = _mock_client([Sample(name="a", value=1)])
        with patch("xlstruct.extraction.engine.build_instructor_client", return_value=client):
            engine = ExtractionEngine(cfg)
            await engine.extract("sheet data", Sample)

        kwargs = client.create_with_completion.await_args.kwargs
        assert kwargs["thinking"] == THINKING_BLOCK
        assert kwargs["temperature"] == 1
        assert kwargs["model"] == "claude-sonnet-4-6"

    async def test_no_thinking_kwargs_when_disabled(self):
        cfg = ExtractorConfig(
            provider="anthropic/claude-sonnet-4-6", thinking=False, temperature=0.0
        )
        client = _mock_client([Sample(name="a", value=1)])
        with patch("xlstruct.extraction.engine.build_instructor_client", return_value=client):
            engine = ExtractionEngine(cfg)
            await engine.extract("sheet data", Sample)

        kwargs = client.create_with_completion.await_args.kwargs
        assert "thinking" not in kwargs
        assert kwargs["temperature"] == 0.0


# * suggest_schema honors thinking (the C1 bug)


class TestSuggestSchemaThinking:
    async def test_thinking_kwargs_reach_llm(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"], ws["B1"] = "Name", "Price"
        ws["A2"], ws["B2"] = "Apple", 1.5
        buf = io.BytesIO()
        wb.save(buf)
        path = tmp_path / "s.xlsx"
        path.write_bytes(buf.getvalue())

        suggested = SuggestedFields(
            model_name="Row",
            fields=[FieldDef(name="name", type="str", nullable=False, description="A")],
        )
        suggest_client = _mock_client(suggested)
        with (
            # ^ Engine built in Extractor.__init__ must not hit a real client
            patch("xlstruct.extraction.engine.build_instructor_client", return_value=MagicMock()),
            patch(
                "xlstruct.extraction.pipeline.build_instructor_client", return_value=suggest_client
            ),
        ):
            extractor = Extractor(provider="anthropic/claude-sonnet-4-6", thinking=True)
            await extractor.suggest_schema(str(path))

        kwargs = suggest_client.create_with_completion.await_args.kwargs
        assert kwargs["thinking"] == THINKING_BLOCK
        assert kwargs["temperature"] == 1


# * B2 — confidence default removed (fail-fast on structural invariant break)


class TestConfidenceFailFast:
    def test_missing_confidence_key_raises(self):
        # ^ A bare record lacks the "{field}_confidence" keys the wrapper guarantees
        with pytest.raises(KeyError):
            _split_confidence([Sample(name="a", value=1)], Sample)


# * C4 — only the provider call is wrapped as ExtractionError; our post-processing is not


class TestExtractionErrorScope:
    async def test_llm_call_error_wrapped_as_extraction_error(self):
        cfg = ExtractorConfig(provider="openai/gpt-4o")
        client = MagicMock()
        client.create_with_completion = AsyncMock(side_effect=RuntimeError("api down"))
        with patch("xlstruct.extraction.engine.build_instructor_client", return_value=client):
            engine = ExtractionEngine(cfg)
            with pytest.raises(ExtractionError) as exc:
                await engine.extract("data", Sample)
        assert exc.value.code == ErrorCode.EXTRACTION_LLM_FAILED

    async def test_postprocessing_error_propagates_unwrapped(self):
        # ^ A bug in our own post-processing must NOT be mislabeled as an LLM failure
        cfg = ExtractorConfig(provider="openai/gpt-4o")
        client = _mock_client([Sample(name="a", value=1)])
        with patch("xlstruct.extraction.engine.build_instructor_client", return_value=client):
            engine = ExtractionEngine(cfg)
            with (
                patch.object(
                    ExtractionEngine,
                    "_split_provenance",
                    side_effect=RuntimeError("post-proc boom"),
                ),
                pytest.raises(RuntimeError, match="post-proc boom"),
            ):
                await engine.extract("data", Sample, track_provenance=True)
